import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

doc = Document()

# Set standard margins (1 inch)
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

def set_font(run, name="Arial", size_pt=11, bold=False, italic=False, color_rgb=(0,0,0)):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def add_h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, name="Arial", size_pt=14, bold=True, color_rgb=(0,0,0))
    return p

def add_h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, name="Arial", size_pt=12, bold=True, color_rgb=(0,0,0))
    return p

def add_h3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, name="Arial", size_pt=11, bold=True, color_rgb=(0,0,0))
    return p

def add_body(text, space_after=4, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_font(r_pre, name="Arial", size_pt=11, bold=True, color_rgb=(0,0,0))
    run = p.add_run(text)
    set_font(run, name="Arial", size_pt=11, color_rgb=(0,0,0))
    return p

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_font(r_pre, name="Arial", size_pt=11, bold=True, color_rgb=(0,0,0))
    run = p.add_run(text)
    set_font(run, name="Arial", size_pt=11, color_rgb=(0,0,0))
    return p

def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'''
        <w:tcBorders {}
            w:top="single" w:top-val="single" w:top-sz="4" w:top-space="0" w:top-color="CCCCCC"
            w:left="single" w:left-val="single" w:left-sz="4" w:left-space="0" w:left-color="CCCCCC"
            w:bottom="single" w:bottom-val="single" w:bottom-sz="4" w:bottom-space="0" w:bottom-color="CCCCCC"
            w:right="single" w:right-val="single" w:right-sz="4" w:right-space="0" w:right-color="CCCCCC"/>
    '''.format(nsdecls('w')))
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex="F4F4F5"):
    shading = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), color_hex))
    cell._tc.get_or_add_tcPr().append(shading)

def add_image_box(img_path, caption, placeholder_text=None, width_inch=5.8):
    if img_path and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        doc.add_picture(img_path, width=Inches(width_inch))
    else:
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.cell(0, 0)
        c.width = Inches(5.8)
        set_cell_border(c)
        set_cell_shading(c, "FAFAFA")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(36)
        p.paragraph_format.space_after = Pt(36)
        r = p.add_run(f"[ {placeholder_text or 'TEMPELKAN GAMBAR DI SINI'} ]\n(Klik di dalam kotak lalu Paste gambar / Insert Picture)")
        set_font(r, name="Arial", size_pt=10, bold=True, color_rgb=(100,100,100))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption)
    set_font(r_cap, name="Arial", size_pt=9.5, italic=True, color_rgb=(60,60,60))

# TITLE
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(12)
title_p.paragraph_format.space_after = Pt(4)
r_title = title_p.add_run("SPESIFIKASI KEBUTUHAN PERANGKAT LUNAK (SRS)")
set_font(r_title, name="Arial", size_pt=16, bold=True, color_rgb=(0,0,0))

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_after = Pt(14)
r_sub = sub_p.add_run("SISTEM INFORMASI POINT OF SALE (POS) CAFE RJ\nDigitalisasi Transaksi Kasir, Multi-Payment, dan Rekapitulasi Penjualan Berbasis Web (Laravel 12 & MySQL)")
set_font(r_sub, name="Arial", size_pt=11, bold=True, color_rgb=(0,0,0))

# Metadata Box
meta_table = doc.add_table(rows=5, cols=2)
meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ("Mata Pelajaran", ": Rekayasa Perangkat Lunak (RPL) / Web Programming"),
    ("Topik Penugasan", ": Rancang Bangun Sistem POS Cafe Berbasis Web"),
    ("Jurusan / Kelas", ": Pengembangan Perangkat Lunak dan GIM (PPLG)"),
    ("Kelompok Siswa", ": 1. [Nama Anggota 1 - Ketua Kelompok]\n  2. [Nama Anggota 2]\n  3. [Nama Anggota 3]\n  4. [Nama Anggota 4]"),
    ("Sekolah / Tahun", ": SMK [Nama Sekolah Anda] / 2026-2027")
]
for i, (k, v) in enumerate(meta_data):
    cell_k = meta_table.cell(i, 0)
    cell_v = meta_table.cell(i, 1)
    cell_k.width = Inches(1.8)
    cell_v.width = Inches(4.7)
    rk = cell_k.paragraphs[0].add_run(k)
    set_font(rk, name="Arial", size_pt=10, bold=True, color_rgb=(0,0,0))
    rv = cell_v.paragraphs[0].add_run(v)
    set_font(rv, name="Arial", size_pt=10, color_rgb=(0,0,0))
    set_cell_border(cell_k)
    set_cell_border(cell_v)

doc.add_page_break()

# LEMBAR PENGESAHAN
add_h1("LEMBAR PENGESAHAN PROYEK")
add_body("Tugas Proyek Rancang Bangun Perangkat Lunak dengan judul Sistem Informasi Point of Sale (POS) Cafe RJ ini telah diselesaikan, diuji, dan disahkan untuk memenuhi penilaian mata pelajaran Rekayasa Perangkat Lunak (RPL).")
add_body("Disahkan di: [Kota]\nPada tanggal: ........................................ 2026", space_after=18)

sign_table = doc.add_table(rows=1, cols=2)
sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cell_g1 = sign_table.cell(0, 0)
cell_g2 = sign_table.cell(0, 1)
cell_g1.width = Inches(3.2)
cell_g2.width = Inches(3.2)

p_g1 = cell_g1.paragraphs[0]
p_g1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_g1.add_run("Mengetahui,\nGuru Pembimbing / Penguji\n\n\n\n___________________________________\nNIP. ........................................")
set_font(r, name="Arial", size_pt=10, bold=True, color_rgb=(0,0,0))

p_g2 = cell_g2.paragraphs[0]
p_g2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_g2.add_run("Disusun Oleh,\nKetua Kelompok / Siswa\n\n\n\n___________________________________\nNIS. ........................................")
set_font(r, name="Arial", size_pt=10, bold=True, color_rgb=(0,0,0))

set_cell_border(cell_g1)
set_cell_border(cell_g2)

doc.add_page_break()

# BAB I
add_h1("BAB I: PENDAHULUAN")
add_h2("A. Latar Belakang Masalah")
add_bullet("Pencatatan nota fisik kasir berisiko hilang, robek, atau basah terkena tumpahan minuman.", "1. Pencatatan Transaksi Masih Manual: ")
add_bullet("Pada jam sibuk (rush hour), kasir kesulitan menghitung subtotal dan uang kembalian secara cepat dengan kalkulator manual, sehingga rawan memicu selisih kas fisik.", "2. Antrean Menumpuk & Rawan Salah Hitung: ")
add_bullet("Pelanggan semakin banyak bertransaksi menggunakan QRIS dan transfer perbankan, sehingga membutuhkan integrasi nominal tagihan langsung pada layar kasir.", "3. Kebutuhan Transaksi Non-Tunai: ")
add_bullet("Pemilik cafe membutuhkan waktu lama saat tutup toko untuk merekapitulasi nota kertas satu per satu guna mengetahui omzet harian.", "4. Rekapitulasi Omzet Lambat: ")
add_bullet("Diperlukan aplikasi kasir web yang ringan, mudah diakses lewat komputer kasir maupun smartphone, serta stabil tanpa kendala sesi kadaluarsa.", "5. Kebutuhan Sistem Web Responsif: ")

add_h2("B. Rumusan Masalah")
add_bullet("Bagaimana merancang dan membangun aplikasi kasir digital (POS) berbasis web dengan Laravel 12 dan MySQL yang mudah digunakan oleh kasir?")
add_bullet("Bagaimana mengintegrasikan sistem kasir dengan tiga metode pembayaran (Tunai dengan kalkulator kembalian otomatis, QRIS dinamis, dan Transfer Bank)?")
add_bullet("Bagaimana mengotomatisasi pencatatan transaksi, pengurangan stok produk, dan pencetakan struk belanja secara instan?")

add_h2("C. Batasan Masalah")
add_bullet("Aplikasi dibangun berbasis web menggunakan framework Laravel 12 dan basis data MySQL.")
add_bullet("Hak akses sistem dibatasi menjadi 2 peran (multi-role): Admin (Pemilik/Manajer Cafe) dan Kasir.")
add_bullet("Metode pembayaran mencakup 3 opsi: Tunai (Cash), QRIS (tampilan standee dan modal pembesar layar pelanggan), serta Transfer Bank manual (BCA, Mandiri, BRI).")
add_bullet("Output cetak struk mencakup format Faktur Layar Penuh (A4) dan format Struk Mini Thermal (80mm).")

add_h2("D. Tujuan Pengembangan")
add_bullet("Menghasilkan aplikasi kasir POS Cafe berbasis web yang fungsional, cepat, dan responsif.")
add_bullet("Membantu kasir melayani pesanan dan menghitung uang kembalian secara otomatis dan akurat.")
add_bullet("Membantu pemilik cafe memantau riwayat transaksi dan laporan omzet penjualan kapan saja.")

add_h2("E. Nama dan Deskripsi Aplikasi")
add_body("Nama aplikasi ini adalah Sistem Informasi Point of Sale Cafe RJ (Cafe RJ POS). Aplikasi ini dirancang untuk mendigitalkan seluruh alur operasional kasir mulai dari pemilihan menu, keranjang pesanan, perhitungan subtotal otomatis, proses pembayaran, hingga cetak struk dan rekapitulasi penjualan.")

add_h2("F. Definisi Istilah dan Singkatan")
add_h3("1. Tabel Definisi Istilah")
tbl_terms = doc.add_table(rows=1, cols=3)
tbl_terms.alignment = WD_TABLE_ALIGNMENT.CENTER
for j, h in enumerate(["No", "Istilah", "Pengertian"]):
    c = tbl_terms.cell(0, j)
    set_cell_shading(c, "E4E4E7")
    set_cell_border(c)
    p = c.paragraphs[0]
    r = p.add_run(h)
    set_font(r, name="Arial", size_pt=10, bold=True, color_rgb=(0,0,0))

terms_data = [
    ("1", "Point of Sale (POS)", "Sistem kasir berbasis software untuk melayani transaksi pembelian pelanggan."),
    ("2", "Admin", "Pengguna yang memiliki hak akses penuh mengelola produk, kategori, user, dan laporan omzet."),
    ("3", "Kasir", "Pengguna yang bertugas mengoperasikan terminal kasir, melayani pesanan, dan mencetak struk."),
    ("4", "Database", "Tempat penyimpanan terpusat data produk, kategori, user, dan transaksi (MySQL)."),
    ("5", "Subtotal", "Nilai total harga dari perkalian kuantitas item yang dibeli dengan harga satuan."),
    ("6", "Uang Kembalian", "Selisih lebih antara uang tunai yang diserahkan pembeli dengan total belanja."),
    ("7", "Struk / Invoice", "Tanda bukti pembayaran resmi yang memuat rincian transaksi belanja.")
]
for row in terms_data:
    r_cells = tbl_terms.add_row().cells
    for j, val in enumerate(row):
        set_cell_border(r_cells[j])
        p = r_cells[j].paragraphs[0]
        r = p.add_run(val)
        set_font(r, name="Arial", size_pt=9.5, color_rgb=(0,0,0))
tbl_terms.columns[0].width = Inches(0.5)
tbl_terms.columns[1].width = Inches(2.0)
tbl_terms.columns[2].width = Inches(4.0)

doc.add_page_break()

# BAB II
add_h1("BAB II: LANDASAN TEORI")
add_h2("A. Konsep Point of Sale (POS) Kasir Kedai Kopi")
add_body("Sistem Point of Sale (POS) adalah peranti utama dalam bisnis makanan dan minuman (F&B) untuk mencatat transaksi penjualan. Sistem POS modern menggantikan mesin kasir konvensional dengan keunggulan:")
add_bullet("Katalog menu visual berbasis kategori untuk mempercepat kasir mencari produk saat melayani antrean.")
add_bullet("Kalkulasi total tagihan dan uang kembalian secara otomatis tanpa perlu kalkulator manual.")
add_bullet("Otomatisasi pemotongan stok bahan baku/menu secara real-time setiap kali transaksi berhasil dibayarkan.")

add_h2("B. Konsep Arsitektur Web Model-View-Controller (MVC)")
add_body("Aplikasi Cafe RJ POS dibangun menggunakan pola desain Model-View-Controller (MVC) bawaan framework Laravel:")
add_bullet("Model: Mengatur struktur data dan relasi basis data (Category, Product, User, Transaction, TransactionDetail).")
add_bullet("View: Mengatur tampilan antarmuka kasir berbasis Blade Templating Engine (HTML dan CSS).")
add_bullet("Controller: Mengatur logika bisnis, memproses masukan kasir, dan menghubungkan Model dengan View (PosController, ProductController, AuthController).")

add_h2("C. Alat Bantu Perancangan Sistem")
add_bullet("Unified Modeling Language (UML): Bahasa pemodelan grafis untuk merancang alur sistem (Use Case Diagram dan Activity Diagram).")
add_bullet("Entity Relationship Diagram (ERD): Diagram pemodelan struktur tabel basis data relasional beserta kunci primer (PK) dan kunci asing (FK).")
add_bullet("Data Flow Diagram (DFD): Diagram aliran pertukaran data antara entitas luar dengan sistem kasir.")

add_h2("D. Teknologi Pengembangan")
add_bullet("Laravel 12: Framework backend PHP modern dengan sistem keamanan bawaan dan ORM Eloquent.")
add_bullet("PHP 8.2+: Bahasa pemrograman backend utama dengan performa tinggi.")
add_bullet("MySQL 8.x: Sistem manajemen basis data relasional (RDBMS) bertipe InnoDB untuk menjamin konsistensi data.")
add_bullet("Tailwind CSS: Framework styling utilitas untuk menciptakan antarmuka kasir yang rapi dan responsif.")
add_bullet("Visual Studio Code / Antigravity IDE: Lingkungan pengembangan terpadu untuk penulisan dan pengujian kode.")

doc.add_page_break()

# BAB III
add_h1("BAB III: METODE DAN PERANCANGAN SISTEM")
add_h2("A. Metode Pengembangan (Model Waterfall)")
add_body("Pengembangan aplikasi menggunakan model sekuensial Waterfall yang terdiri dari 5 tahapan:")
add_bullet("Analisis Kebutuhan: Mengidentifikasi kebutuhan fungsional kasir cafe dan spesifikasi fitur MVP.")
add_bullet("Perancangan Sistem: Membuat diagram Use Case, Activity Diagram, ERD, DFD, serta wireframe dan mockup antarmuka.")
add_bullet("Pengodean (Implementation): Menulis kode backend MVC Laravel 12 dan template frontend kasir.")
add_bullet("Pengujian (Testing): Menguji seluruh fungsionalitas dengan metode Black-Box Testing.")
add_bullet("Pemeliharaan (Maintenance): Mengoptimalkan responsivitas antarmuka di PC kasir maupun smartphone.")

add_h2("B. Analisis Kebutuhan Pengguna (Actors)")
add_h3("1. Kebutuhan Role Kasir:")
add_bullet("Melakukan login akun kasir untuk memulai shift kerja.")
add_bullet("Memilih menu berdasarkan kategori (Semua Menu, Kopi, Non-Kopi, Makanan Ringan).")
add_bullet("Mengatur kuantitas item dalam keranjang pesanan.")
add_bullet("Memilih metode pembayaran: Tunai (hitung kembalian), QRIS (tampilan barcode), atau Transfer Bank.")
add_bullet("Menyelesaikan pembayaran dan mencetak struk belanja (format Faktur A4 dan Struk Thermal 80mm).")
add_bullet("Melihat daftar riwayat transaksi kasir.")

add_h3("2. Kebutuhan Role Admin:")
add_bullet("Melakukan login akun admin dengan hak akses penuh.")
add_bullet("Mengelola master kategori menu (tambah, edit, hapus).")
add_bullet("Mengelola master produk (tambah, edit, hapus, atur harga, stok, dan upload foto produk).")
add_bullet("Mengelola data akun pengguna (kasir dan admin).")
add_bullet("Melihat laporan riwayat transaksi dan rekapitulasi omzet penjualan dengan filter tanggal.")

add_h2("C. Pemodelan Sistem (UML & Aliran Data)")

# 1. USE CASE
add_h3("1. Use Case Diagram")
add_body("Diagram Use Case di bawah ini menggambarkan interaksi antara Aktor Kasir dan Aktor Admin dengan seluruh fungsi yang tersedia pada Aplikasi Kasir Cafe RJ:")
add_image_box("docs/use_case_diagram.png", "Gambar 3.1: Use Case Diagram Aplikasi Kasir Cafe RJ", "DIAGRAM USE CASE")

add_body("Tabel Pemetaan Hak Akses Fitur Antara Kasir dan Admin:")
tbl_uc = doc.add_table(rows=1, cols=3)
tbl_uc.alignment = WD_TABLE_ALIGNMENT.CENTER
for j, h in enumerate(["No. Fitur / Use Case", "Kasir", "Admin"]):
    c = tbl_uc.cell(0, j)
    set_cell_shading(c, "E4E4E7")
    set_cell_border(c)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(h)
    set_font(r, name="Arial", size_pt=10, bold=True, color_rgb=(0,0,0))

uc_rows = [
    ("1. Login Akun", "Ya", "Ya"),
    ("2. Pilih Menu & Kategori", "Ya", "Ya"),
    ("3. Kelola Keranjang Pesanan", "Ya", "Ya"),
    ("4. Hitung Bayar & Kembalian", "Ya", "Ya"),
    ("5. Cetak Struk Pembelian", "Ya", "Ya"),
    ("6. Lihat Riwayat Kasir", "Ya", "Ya"),
    ("7. Kelola Kategori Menu (CRUD)", "Tidak", "Ya"),
    ("8. Kelola Data Produk & Foto (CRUD)", "Tidak", "Ya"),
    ("9. Kelola Akun Kasir (CRUD)", "Tidak", "Ya"),
    ("10. Lihat Laporan Omzet Penjualan", "Tidak", "Ya")
]
for row in uc_rows:
    r_cells = tbl_uc.add_row().cells
    for j, val in enumerate(row):
        set_cell_border(r_cells[j])
        p = r_cells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(val)
        set_font(r, name="Arial", size_pt=9.5, color_rgb=(0,0,0))
tbl_uc.columns[0].width = Inches(3.5)
tbl_uc.columns[1].width = Inches(1.5)
tbl_uc.columns[2].width = Inches(1.5)

doc.add_page_break()

# 2. ACTIVITY DIAGRAM
add_h3("2. Activity Diagram (Alur Transaksi Kasir)")
add_body("Activity Diagram berikut menggambarkan urutan langkah prosedural kasir dalam melayani pesanan, memilih metode bayar, hingga mencetak bukti struk:")
add_image_box("docs/activity_diagram.png", "Gambar 3.2: Activity Diagram Alur Transaksi Kasir POS", "DIAGRAM ACTIVITY TRANSAKSI KASIR", width_inch=4.8)

doc.add_page_break()

# 3. ERD
add_h3("3. Entity Relationship Diagram (ERD)")
add_body("ERD di bawah ini memetakan relasi antara 5 entitas tabel basis data MySQL pada database caferj:")
add_image_box("docs/erd_diagram.png", "Gambar 3.3: Entity Relationship Diagram (ERD) Database caferj", "DIAGRAM ERD DATABASE")

# 4. DFD
add_h3("4. Data Flow Diagram (DFD) Konteks Level 0")
add_body("DFD Level 0 berikut memodelkan aliran data masukan dan keluaran antara entitas luar (Kasir dan Admin) dengan sistem POS Cafe RJ:")
add_image_box("docs/dfd_level_0.png", "Gambar 3.4: DFD Konteks Level 0 Sistem POS Cafe RJ", "DIAGRAM DFD KONTEKS LEVEL 0", width_inch=5.4)

doc.add_page_break()

# 5. SKENARIO USE CASE
add_h3("5. Tabel Skenario Use Case")

def add_scenario_table(title, actor, goal, steps):
    add_body(f"Aktor: {actor}\nTujuan: {goal}", bold_prefix=f"{title}\n")
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(["No", "Aksi Pengguna (Aktor)", "Respon Sistem"]):
        c = tbl.cell(0, j)
        set_cell_shading(c, "E4E4E7")
        set_cell_border(c)
        p = c.paragraphs[0]
        r = p.add_run(h)
        set_font(r, name="Arial", size_pt=10, bold=True, color_rgb=(0,0,0))
    for s in steps:
        r_cells = tbl.add_row().cells
        for j, val in enumerate(s):
            set_cell_border(r_cells[j])
            p = r_cells[j].paragraphs[0]
            r = p.add_run(val)
            set_font(r, name="Arial", size_pt=9.5, color_rgb=(0,0,0))
    tbl.columns[0].width = Inches(0.5)
    tbl.columns[1].width = Inches(3.0)
    tbl.columns[2].width = Inches(3.0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

add_scenario_table(
    "a. Skenario Use Case: Login Akun",
    "Kasir / Admin",
    "Masuk ke dalam sistem sesuai hak akses.",
    [
        ("1", "Membuka halaman web kasir pada URL /login.", "Menampilkan formulir login (email dan password)."),
        ("2", "Memasukkan email dan password valid, lalu klik Login.", "Memvalidasi kredensial pengguna."),
        ("3", "Sistem mengarahkan pengguna ke halaman yang sesuai.", "Jika kasir: masuk ke terminal POS. Jika admin: masuk ke dashboard manajemen menu."),
        ("Alt", "Memasukkan email atau password salah.", "Menampilkan pesan peringatan: 'Kredensial tidak valid' dan tetap di halaman login.")
    ]
)

add_scenario_table(
    "b. Skenario Use Case: Transaksi Kasir POS",
    "Kasir",
    "Memilih menu pesanan, menerima pembayaran, dan mencetak struk belanja.",
    [
        ("1", "Memilih kategori menu dan mengklik item produk.", "Menambahkan item ke keranjang dan mengalkulasi subtotal secara real-time."),
        ("2", "Mengatur jumlah kuantitas pesanan (+ atau -).", "Memperbarui total belanja belanja secara otomatis."),
        ("3", "Memilih metode bayar (Tunai, QRIS, atau Transfer).", "Menampilkan panel metode pembayaran yang dipilih."),
        ("4", "Untuk Tunai: Memasukkan nominal uang diterima.", "Sistem otomatis menghitung nilai uang kembalian."),
        ("5", "Kasir menekan tombol 'Bayar & Cetak Struk'.", "Sistem memvalidasi transaksi, menyimpan ke database, mengurangi stok produk, dan membuka jendela cetak struk.")
    ]
)

add_scenario_table(
    "c. Skenario Use Case: Kelola Produk dan Upload Foto",
    "Admin",
    "Menambahkan atau memperbarui data menu dan foto produk cafe.",
    [
        ("1", "Mengklik menu Tambah Produk pada panel admin.", "Menampilkan form nama menu, kategori, harga, stok, dan upload foto."),
        ("2", "Memilih file foto dari komputer.", "Menampilkan live preview gambar menu sebelum disimpan."),
        ("3", "Menekan tombol Simpan Produk.", "Sistem memvalidasi data, menyimpan berkas gambar ke storage public, dan menyimpan record ke database MySQL.")
    ]
)

doc.add_page_break()

# D. PERANCANGAN MOCKUP / WIREFRAME (HITAM PUTIH - LOW FIDELITY) -> 3 FOTO HITAM PUTIH
add_h2("D. Perancangan Mockup / Wireframe (Hitam Putih - Low Fidelity)")
add_body("Bagian ini menyajikan rancangan kerangka kasar dasar antarmuka (wireframe monokrom hitam putih) untuk memetakan tata letak fungsional komponen sebelum diberi desain warna dan styling visual akhir. Terdapat 3 gambar mockup hitam putih:")

# Mockup 1 Hitam Putih: Login
add_h3("1. Mockup / Wireframe Hitam Putih 1: Halaman Login Pengguna")
add_body("Memetakan tata letak kotak formulir autentikasi akun kasir dan admin (input email, password, dan tombol masuk):")
add_image_box("docs/wireframe_login.png", "Gambar 3.5: Wireframe Hitam Putih 1 - Halaman Login Pengguna", "MOCKUP 1: WIREFRAME LOGIN HITAM PUTIH", width_inch=5.4)

# Mockup 2 Hitam Putih: Kasir POS
add_h3("2. Mockup / Wireframe Hitam Putih 2: Terminal Kasir POS Utama")
add_body("Memetakan tata letak bilah kategori menu, susunan grid kartu produk, serta panel keranjang belanja dan tombol checkout:")
add_image_box("docs/wireframe_kasir_pos.png", "Gambar 3.6: Wireframe Hitam Putih 2 - Terminal Kasir POS Utama", "MOCKUP 2: WIREFRAME KASIR POS HITAM PUTIH", width_inch=5.8)

# Mockup 3 Hitam Putih: Admin Dashboard
add_h3("3. Mockup / Wireframe Hitam Putih 3: Panel Administrator dan Master Data")
add_body("Memetakan tata letak sidebar menu admin, 4 kartu ringkasan omzet, serta tabel master data produk cafe:")
add_image_box("docs/wireframe_admin_dashboard.png", "Gambar 3.7: Wireframe Hitam Putih 3 - Panel Administrator & Master Data", "MOCKUP 3: WIREFRAME PANEL ADMIN HITAM PUTIH", width_inch=5.8)

doc.add_page_break()

# E. PERANCANGAN UI/UX & HASIL TAMPILAN WEB (FULL COLOR - HIGH FIDELITY) -> 3 FOTO UI/UX
add_h2("E. Perancangan UI/UX & Hasil Tampilan Web (Full Color - High Fidelity)")
add_body("Bagian ini menyajikan hasil implementasi antarmuka pengguna (UI/UX) akhir yang sudah berwarna penuh (full color), berfoto produk asli, dan aktif beroperasi pada sistem web Cafe RJ POS. Terdapat 3 gambar tampilan UI/UX web:")

# UI/UX 1 Web Color: Login
add_h3("1. UI/UX Hasil Tampilan Web 1: Halaman Login Sistem POS")
add_body("Tampilan antarmuka login aktif di browser web dengan kartu terpusat, logo Cafe RJ, dan form autentikasi multi-role:")
add_image_box("docs/login_web_screen.png", "Gambar 3.8: UI/UX Tampilan Web 1 - Halaman Login Sistem POS (Full Color)", "UI/UX 1: HASIL TAMPILAN WEB HALAMAN LOGIN", width_inch=5.4)

# UI/UX 2 Web Color: Kasir POS
add_h3("2. UI/UX Hasil Tampilan Web 2: Terminal Kasir POS Aktif")
add_body("Tampilan kasir aktif di browser web dengan nuansa gelap (dark zinc), katalog foto menu berharga IDR, stepper keranjang, dan tab pembayaran multi-metode:")
add_image_box("docs/pos_cashier_mockup.jpg", "Gambar 3.9: UI/UX Tampilan Web 2 - Terminal Kasir POS Aktif (Full Color)", "UI/UX 2: HASIL TAMPILAN WEB KASIR POS", width_inch=5.8)

# UI/UX 3 Web Color: Admin Dashboard
add_h3("3. UI/UX Hasil Tampilan Web 3: Dashboard & Manajemen Produk Admin")
add_body("Tampilan panel admin aktif di browser web dengan metrik penjualan real-time serta tabel CRUD master menu lengkap dengan thumbnail gambar:")
add_image_box("docs/admin_dashboard_mockup.jpg", "Gambar 3.10: UI/UX Tampilan Web 3 - Dashboard & Manajemen Menu Admin (Full Color)", "UI/UX 3: HASIL TAMPILAN WEB DASHBOARD ADMIN", width_inch=5.8)

doc.add_page_break()

# BAB IV
add_h1("BAB IV: LEMBAR EVALUASI DAN PENILAIAN GURU")
add_body("Lembar evaluasi ini digunakan oleh guru pembimbing / penguji untuk memberikan nilai terhadap hasil rancang bangun perangkat lunak yang diselesaikan oleh kelompok siswa.")

add_h2("A. Rubrik Penilaian Proyek")
tbl_rubrik = doc.add_table(rows=1, cols=5)
tbl_rubrik.alignment = WD_TABLE_ALIGNMENT.CENTER
headers_rubrik = ["No", "Aspek Penilaian", "Bobot", "Skor (1-100)", "Catatan Penguji"]
for j, h in enumerate(headers_rubrik):
    c = tbl_rubrik.cell(0, j)
    set_cell_shading(c, "E4E4E7")
    set_cell_border(c)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [0, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(h)
    set_font(r, name="Arial", size_pt=10, bold=True, color_rgb=(0,0,0))

rubrik_rows = [
    ("1", "Fungsionalitas Sistem (40%):\n- CRUD kategori, produk, dan user berjalan lancar\n- Login multi-role sesuai hak akses\n- Hitungan kasir dan uang kembalian akurat\n- Cetak struk berjalan tanpa kendala", "40%", "", ""),
    ("2", "Kerapian Struktur Kode (20%):\n- Penerapan arsitektur MVC Laravel 12 yang tepat\n- Kerapian penataan route, controller, model, dan view\n- Bersih dari kode error", "20%", "", ""),
    ("3", "Desain UI/UX & Responsivitas (20%):\n- Antarmuka kasir intuitif, mudah dipahami, dan cepat\n- Tampilan responsif di layar komputer maupun handphone\n- Tata letak rapi dan estetis", "20%", "", ""),
    ("4", "Kelengkapan Basis Data (20%):\n- Desain tabel MySQL efisien dan ternormalisasi\n- Relasi foreign key antar tabel berjalan tepat\n- Seeder dan migration lengkap", "20%", "", ""),
    ("", "TOTAL NILAI AKHIR (100%):", "100%", "", "Predikat: ....................")
]
for row in rubrik_rows:
    r_cells = tbl_rubrik.add_row().cells
    for j, val in enumerate(row):
        set_cell_border(r_cells[j])
        p = r_cells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [0, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(val)
        is_bold = True if row == rubrik_rows[-1] else False
        set_font(r, name="Arial", size_pt=9.5, bold=is_bold, color_rgb=(0,0,0))

tbl_rubrik.columns[0].width = Inches(0.4)
tbl_rubrik.columns[1].width = Inches(2.8)
tbl_rubrik.columns[2].width = Inches(0.8)
tbl_rubrik.columns[3].width = Inches(1.1)
tbl_rubrik.columns[4].width = Inches(1.4)

add_h2("B. Lembar Catatan dan Feedback Penguji")
add_body("Catatan / Masukan Perbaikan dari Guru Penguji:")
add_body("........................................................................................................................................................................")
add_body("........................................................................................................................................................................")
add_body("........................................................................................................................................................................", space_after=20)

sign_eval = doc.add_table(rows=1, cols=2)
sign_eval.alignment = WD_TABLE_ALIGNMENT.CENTER
c1 = sign_eval.cell(0, 0)
c2 = sign_eval.cell(0, 1)
c1.width = Inches(3.2)
c2.width = Inches(3.2)

p1 = c1.paragraphs[0]
p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p1.add_run("Guru Penguji / Penilai 1\n\n\n\n___________________________________\nNIP. ........................................")
set_font(r, name="Arial", size_pt=10, bold=True, color_rgb=(0,0,0))

p2 = c2.paragraphs[0]
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p2.add_run("Guru Penguji / Penilai 2\n\n\n\n___________________________________\nNIP. ........................................")
set_font(r, name="Arial", size_pt=10, bold=True, color_rgb=(0,0,0))

set_cell_border(c1)
set_cell_border(c2)

doc.save("docs/SRS_Cafe_RJ.docx")
doc.save("SRS_Cafe_RJ.docx")
print("SUCCESS: SRS_Cafe_RJ.docx regenerated with exactly 3 Mockups Hitam Putih + 3 UI/UX Web!")
